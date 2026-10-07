import openpyxl
from app.config.settings import EXCEL_SOURCE
from app.auditor.models import AuditTask
import logging

logger = logging.getLogger(__name__)

class ExcelLoader:
    def __init__(self, file_path=EXCEL_SOURCE):
        self.file_path = file_path

    def load_tasks(self) -> list[AuditTask]:
        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=False)
            all_tasks = []
            
            aliases = {
                "Mercado Libre": ["Meli", "Mercado Libre", "Mercadolibre"],
                "Paris": ["Paris", "Paris.cl"],
                "Ripley": ["Ripley", "Ripley.cl"],
                "Falabella": ["Falabella", "Falabella.com"]
            }

            for sheet_name in wb.sheetnames:
                marketplace_raw = sheet_name.split("(")[0].strip()
                ws = wb[sheet_name]
                if ws.max_row < 2:
                    continue

                headers = {}
                for col_idx in range(1, ws.max_column + 1):
                    val = ws.cell(row=1, column=col_idx).value
                    if val:
                        headers[col_idx] = str(val).strip()

                search_terms = aliases.get(marketplace_raw, [marketplace_raw])

                cat_mkt_col = None
                cat_nico_col = None

                for col_idx, h in headers.items():
                    if "Categor" in h:
                        if "Nicopoly" in h:
                            cat_nico_col = col_idx
                        elif any(term.lower() in h.lower() for term in search_terms) or not cat_mkt_col:
                            cat_mkt_col = col_idx

                if not cat_mkt_col or not cat_nico_col:
                    logger.warning(f"Could not find category columns in sheet {sheet_name}. Available: {list(headers.values())}")
                    continue

                for row_idx in range(2, ws.max_row + 1):
                    cell_mkt = ws.cell(row=row_idx, column=cat_mkt_col)
                    cell_nico = ws.cell(row=row_idx, column=cat_nico_col)

                    val_mkt = cell_mkt.value
                    if val_mkt is None or str(val_mkt).strip() == "":
                        continue

                    category_marketplace = str(val_mkt).strip()
                    if marketplace_raw.lower() == "ripley" and category_marketplace.lower() == "accesorios y complemento":
                        category_marketplace = "Accesorios y complementos"
                    val_nico = cell_nico.value if cell_nico else None
                    category_nicopoly = str(val_nico).strip() if val_nico else category_marketplace

                    url_base = cell_mkt.hyperlink.target if cell_mkt.hyperlink else None
                    if not url_base and category_marketplace.startswith("http"):
                        url_base = category_marketplace

                    url_nicopoly = cell_nico.hyperlink.target if cell_nico.hyperlink else None
                    if not url_nicopoly and category_nicopoly.startswith("http"):
                        url_nicopoly = category_nicopoly

                    task = AuditTask(
                        marketplace=marketplace_raw,
                        category_marketplace=category_marketplace,
                        category_nicopoly=category_nicopoly,
                        url_base=url_base,
                        url_nicopoly=url_nicopoly
                    )
                    all_tasks.append(task)

            # De-duplicate tasks based on marketplace and category_marketplace
            # Keep valid URLs and canonical category naming across historical rows
            unique_tasks = {}
            for task in all_tasks:
                key = (task.marketplace.lower(), task.category_marketplace.lower())
                if key not in unique_tasks:
                    unique_tasks[key] = task
                else:
                    existing = unique_tasks[key]
                    if task.url_base:
                        existing.url_base = task.url_base
                    if task.url_nicopoly:
                        existing.url_nicopoly = task.url_nicopoly
                    existing.category_marketplace = task.category_marketplace
                    existing.category_nicopoly = task.category_nicopoly

            return list(unique_tasks.values())
        except Exception as e:
            logger.error(f"Error loading tasks from Excel: {e}")
            return []

