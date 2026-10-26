"""Season CHECKLIST reference (validation-only, never acquisition input).

The checklist provides per-channel IDs/links for the current season. It is
consumed AFTER autonomous acquisition to verify publication coverage; it must
never seed queries, categories or surfaces.
"""
from __future__ import annotations

import json
import os
import re
import unicodedata
from datetime import datetime
from pathlib import Path

CHANNELS = {
    "falabella": ("Falabella", "Falabella_Link", "Falabella_Observ."),
    "ripley": ("MPM_RIPLEY", "RIPLEY_LINK", "RIPLEY_OBSERV."),
    "paris": ("MKP_PARIS", "PARIS_LINK", "PARIS_OBSERV."),
    "mercado libre": ("MLC_MELI", "MELI_LINK", "MELI_OBSERV."),
}
CACHE_NAME = "checklist_cache.json"


def default_paths():
    out = []
    env = os.environ.get("MRI_CHECKLIST_PATH")
    if env:
        out.append(Path(env))
    root = Path(__file__).resolve().parents[2]
    out.append(root / "CHECKLIST (GRUPOS_2026 SS).xlsx")
    out.append(Path.home() / "Downloads" / "CHECKLIST (GRUPOS_2026 SS).xlsx")
    return out


def find_checklist():
    for p in default_paths():
        if p and p.exists():
            return p
    return None


def _clean_id(value):
    if value is None:
        return None
    s = str(value).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s or None


def _url_product_id(url):
    m = re.search(r"/product/(\d+)", str(url or ""))
    return m.group(1) if m else None


def _norm(text):
    s = unicodedata.normalize("NFKD", str(text or "")).encode(
        "ascii", "ignore").decode().lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", s).split())


def load_checklist(path=None, *, use_cache=True):
    path = Path(path) if path else find_checklist()
    if not path or not path.exists():
        return None
    cache = path.parent / CACHE_NAME
    try:
        if use_cache and cache.exists() and cache.stat().st_mtime >= path.stat().st_mtime:
            return json.loads(cache.read_text(encoding="utf-8"))
    except Exception:
        pass
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    products = []
    try:
        for ws in wb.worksheets:
            header_row, cols = None, {}
            for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=8,
                                                     values_only=True), start=1):
                if any(str(c).strip().upper() == "SKU_PADRE" for c in row if c):
                    header_row = r_idx
                    cols = {str(c).strip(): i for i, c in enumerate(row) if c}
                    break
            if not header_row:
                continue
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                sku_padre = _clean_id(row[cols["SKU_PADRE"]]) if "SKU_PADRE" in cols else None
                sku_hijo = _clean_id(row[cols["SKU_HIJO"]]) if "SKU_HIJO" in cols else None
                name = row[cols["NOMBRE_PRODUCTO"]] if "NOMBRE_PRODUCTO" in cols else None
                if not (sku_padre or sku_hijo) and not name:
                    continue
                channels = {}
                for ch, (id_col, link_col, obs_col) in CHANNELS.items():
                    cid = _clean_id(row[cols[id_col]]) if id_col in cols else None
                    link = row[cols[link_col]] if link_col in cols else None
                    obs = row[cols[obs_col]] if obs_col in cols else None
                    channels[ch] = {"product_id": cid, "link": link,
                                    "observed_manual": obs}
                products.append({"group": ws.title, "sku_padre": sku_padre,
                                 "sku_hijo": sku_hijo, "name": name,
                                 "channels": channels})
    finally:
        wb.close()
    data = {"source": str(path), "generated_at": datetime.now().isoformat(),
            "authority": "VALIDATION_ONLY",
            "products": products}
    try:
        cache.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    except Exception:
        pass
    return data


def build_checklist_report(checklist, marketplace, observations, *, run_id=None,
                           run_complete=True):
    channel = marketplace.lower()
    parents = {}
    id_index = {}
    sku_index = {}
    title_index = {}
    for p in (checklist or {}).get("products") or []:
        ch = (p.get("channels") or {}).get(channel) or {}
        pid = p.get("sku_padre") or p.get("sku_hijo")
        if not pid:
            continue
        parents.setdefault(pid, {"parent": pid, "name": p.get("name"),
                                 "observed_manual": ch.get("observed_manual")})
        if ch.get("product_id"):
            id_index[str(ch["product_id"])] = pid
        url_id = _url_product_id(ch.get("link"))
        if url_id:
            id_index[str(url_id)] = pid
        for sku in (p.get("sku_padre"), p.get("sku_hijo")):
            if sku:
                sku_index[str(sku).lower()] = pid
        if _norm(p.get("name")):
            title_index.setdefault(_norm(p.get("name")), pid)
    published, extra = {}, []
    for obs in observations or []:
        text = " ".join(str(obs.get(k) or "") for k in
                        ("producto_id", "sku", "producto_nombre", "canonical_url"))
        pid = None
        method = None
        for key in re.findall(r"\d{6,}", text):
            if key in id_index:
                pid, method = id_index[key], "CHECKLIST_ID"
                break
        if pid is None:
            for sku, cand in sku_index.items():
                if sku and re.search(r"(?<![A-Za-z0-9])" + re.escape(sku)
                                     + r"(?![A-Za-z0-9])", text, re.I):
                    pid, method = cand, "CHECKLIST_SKU"
                    break
        if pid is None:
            n = _norm(obs.get("producto_nombre"))
            pid, method = title_index.get(n), ("CHECKLIST_TITLE" if n else None)
        if pid is None:
            extra.append({"observed_id": obs.get("producto_id"),
                          "title": obs.get("producto_nombre"),
                          "url": obs.get("canonical_url")})
            continue
        if pid not in published:
            published[pid] = {"parent": pid, "name": parents[pid]["name"],
                              "method": method,
                              "observed_id": obs.get("producto_id"),
                              "url": obs.get("canonical_url")}
    missing = [{"parent": pid, "name": g["name"]}
               for pid, g in parents.items() if pid not in published]
    coverage = {
        "published_verified": len(published) if run_complete else 0,
        "missing": len(missing) if run_complete else 0,
        "unverified": 0 if run_complete else len(parents),
        "extra_not_in_checklist": len(extra),
    }
    return {"marketplace": marketplace, "run_id": run_id,
            "generated_at": datetime.now().isoformat(),
            "authority": "VALIDATION_ONLY", "scope": "SEASON_CHECKLIST",
            "run_complete": run_complete,
            "total_products": len(parents), "coverage": coverage,
            "published": list(published.values()), "missing": missing,
            "extra": extra[:500]}
